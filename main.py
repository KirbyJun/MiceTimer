"""
MiceTimer v2.0 - 实验计时计数软件
PySide6 GUI, Windows-compatible
Hierarchical: Paradigm → Groups (对照/实验) → Subjects → Items
"""
from __future__ import annotations

import json
import os
import sys
import time
from copy import deepcopy
from datetime import datetime
from typing import Any, Dict, List, Optional

from PySide6.QtCore import QSize, Qt, QTimer, Signal
from PySide6.QtGui import QColor, QFont, QKeySequence
from PySide6.QtWidgets import (
    QAbstractItemView,
    QApplication,
    QComboBox,
    QDialog,
    QFileDialog,
    QFormLayout,
    QGroupBox,
    QHBoxLayout,
    QHeaderView,
    QInputDialog,
    QKeySequenceEdit,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QScrollArea,
    QSizePolicy,
    QStatusBar,
    QTableWidget,
    QTableWidgetItem,
    QTabWidget,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------

def _app_dir() -> str:
    """Return the directory where the application executable (or script) lives."""
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


APP_DIR = _app_dir()
DATA_DIR = os.path.join(APP_DIR, "data")
AUTOSAVE_DIR = os.path.join(DATA_DIR, "autosave")
EXPORT_DIR = os.path.join(DATA_DIR, "export")
RECOVERY_FILE = os.path.join(AUTOSAVE_DIR, "recovery.json")
SETTINGS_FILE = os.path.join(DATA_DIR, "settings.json")
RECENT_TEMPLATE_FILE = os.path.join(DATA_DIR, "recent_template.json")
TEMPLATES_DIR = os.path.join(DATA_DIR, "templates")

for _d in (AUTOSAVE_DIR, EXPORT_DIR, TEMPLATES_DIR):
    os.makedirs(_d, exist_ok=True)

MAX_FILENAME_PART_LENGTH = 30
MAX_RECENT_TEMPLATES = 10

# ---------------------------------------------------------------------------
# Paradigm / group constants
# ---------------------------------------------------------------------------

PARADIGM_3SIT = "3-SIT"
PARADIGM_FREESIT = "Free-SIT"
ALL_PARADIGMS = [PARADIGM_3SIT, PARADIGM_FREESIT]

GROUP_CONTROL = "对照组"
GROUP_EXPERIMENT = "实验组"
ALL_GROUPS = [GROUP_CONTROL, GROUP_EXPERIMENT]

DEFAULT_ITEMS: Dict[str, List[Dict]] = {
    PARADIGM_3SIT: [
        {"kind": "timer", "name": "Mice/s"},
        {"kind": "timer", "name": "Toy/s"},
    ],
    PARADIGM_FREESIT: [
        {"kind": "timer", "name": "适应时间/s"},
        {"kind": "timer", "name": "嗅探时间/s"},
        {"kind": "counter", "name": "躲避次数"},
    ],
}

DEFAULT_SETTINGS: Dict[str, Any] = {
    "hotkeys": {
        "start_stop": "F5",
        "reset": "F6",
        "export": "F7",
    },
}

# ---------------------------------------------------------------------------
# Format helpers
# ---------------------------------------------------------------------------

def fmt_ssxx(seconds: float) -> str:
    """Format seconds -> SS.xx (truncate to centiseconds)."""
    if seconds < 0:
        seconds = 0.0
    total_cs = int(seconds * 100)
    s = total_cs // 100
    cs = total_cs % 100
    return f"{s}.{cs:02d}"


def fmt_ssxx_signed(seconds: float) -> str:
    """Format possibly-negative seconds -> ±SS.xx."""
    if seconds < 0:
        return f"-{fmt_ssxx(-seconds)}"
    return fmt_ssxx(seconds)


def safe_name_part(s: str) -> str:
    """Make a string safe for use as a filename component."""
    if not s:
        return ""
    return "".join(c for c in s if c.isalnum() or c in "-_ ")[:MAX_FILENAME_PART_LENGTH].strip()


def calc_di(mice_s: float, toy_s: float) -> str:
    """Return DI = (Mice-Toy)/(Mice+Toy) formatted to 6 decimal places, or '' if denominator is 0."""
    mice_plus_toy = mice_s + toy_s
    if mice_plus_toy == 0:
        return ""
    return f"{(mice_s - toy_s) / mice_plus_toy:.6f}"


def calc_di_numeric(mice_s: float, toy_s: float) -> Optional[float]:
    """Return DI = (Mice-Toy)/(Mice+Toy) as float, or None if denominator is 0."""
    total = mice_s + toy_s
    if total == 0:
        return None
    return (mice_s - toy_s) / total


# ---------------------------------------------------------------------------
# Settings helpers
# ---------------------------------------------------------------------------

def load_settings() -> dict:
    if os.path.exists(SETTINGS_FILE):
        try:
            with open(SETTINGS_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
            merged = deepcopy(DEFAULT_SETTINGS)
            for k, v in data.items():
                if k == "hotkeys" and isinstance(v, dict):
                    merged["hotkeys"].update(v)
                else:
                    merged[k] = v
            return merged
        except Exception:
            pass
    return deepcopy(DEFAULT_SETTINGS)


def save_settings(settings: dict):
    with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
        json.dump(settings, f, ensure_ascii=False, indent=2)


def load_recent_templates() -> List[dict]:
    if os.path.exists(RECENT_TEMPLATE_FILE):
        try:
            with open(RECENT_TEMPLATE_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
            if isinstance(data, list):
                return data
        except Exception:
            pass
    return []


def save_recent_template(template: dict):
    recent = load_recent_templates()
    name = template.get("name", "")
    recent = [t for t in recent if t.get("name") != name]
    recent.insert(0, template)
    recent = recent[:MAX_RECENT_TEMPLATES]
    with open(RECENT_TEMPLATE_FILE, "w", encoding="utf-8") as f:
        json.dump(recent, f, ensure_ascii=False, indent=2)


# ---------------------------------------------------------------------------
# Data model
# ---------------------------------------------------------------------------

class Item:
    def __init__(self, kind: str, name: str):
        self.kind: str = kind           # "timer" | "counter"
        self.name: str = name
        self.elapsed: float = 0.0       # seconds (timer)
        self.count: int = 0             # (counter)
        self.running: bool = False      # timer only
        self.last_start_ts: Optional[float] = None

    def to_dict(self) -> dict:
        return {
            "kind": self.kind,
            "name": self.name,
            "elapsed": self.elapsed,
            "count": self.count,
            "running": self.running,
            "last_start_ts": self.last_start_ts,
        }

    @classmethod
    def from_dict(cls, d: dict) -> "Item":
        it = cls(d["kind"], d["name"])
        it.elapsed = float(d.get("elapsed", 0))
        it.count = int(d.get("count", 0))
        it.running = bool(d.get("running", False))
        it.last_start_ts = d.get("last_start_ts", None)
        return it

    def reset(self):
        self.elapsed = 0.0
        self.count = 0
        self.running = False
        self.last_start_ts = None

    def current_elapsed(self) -> float:
        if self.running and self.last_start_ts is not None:
            return self.elapsed + (time.perf_counter() - self.last_start_ts)
        return self.elapsed


class Subject:
    """One experimental subject (e.g., B21) belonging to a group."""

    def __init__(self, name: str, paradigm: str):
        self.name: str = name
        self.paradigm: str = paradigm
        self.items: List[Item] = [
            Item(it["kind"], it["name"])
            for it in DEFAULT_ITEMS.get(paradigm, [])
        ]
        self.started: bool = False
        self.start_time: Optional[str] = None

    def to_dict(self) -> dict:
        return {
            "name": self.name,
            "paradigm": self.paradigm,
            "items": [it.to_dict() for it in self.items],
            "started": self.started,
            "start_time": self.start_time,
        }

    @classmethod
    def from_dict(cls, d: dict) -> "Subject":
        subj = cls.__new__(cls)
        subj.name = d.get("name", "")
        subj.paradigm = d.get("paradigm", PARADIGM_3SIT)
        subj.items = [Item.from_dict(it) for it in d.get("items", [])]
        subj.started = d.get("started", False)
        subj.start_time = d.get("start_time", None)
        return subj

    def reset(self):
        for it in self.items:
            it.reset()
        self.started = False
        self.start_time = None

    def get_value(self, item_name: str) -> float:
        """Return current elapsed seconds or count for item by name."""
        for it in self.items:
            if it.name == item_name:
                if it.kind == "timer":
                    return it.current_elapsed()
                return float(it.count)
        return 0.0


class Group:
    """One experimental group (对照组 or 实验组)."""

    def __init__(self, name: str):
        self.name: str = name
        self.subjects: List[Subject] = []

    def to_dict(self) -> dict:
        return {
            "name": self.name,
            "subjects": [s.to_dict() for s in self.subjects],
        }

    @classmethod
    def from_dict(cls, d: dict) -> "Group":
        g = cls(d.get("name", ""))
        g.subjects = [Subject.from_dict(s) for s in d.get("subjects", [])]
        return g


class Session:
    """Top-level data container for one experiment session."""

    def __init__(self, paradigm: str = PARADIGM_3SIT):
        self.paradigm: str = paradigm
        self.date: str = datetime.now().strftime("%Y-%m-%d")
        self.operator: str = ""
        self.remark: str = ""
        self.groups: List[Group] = [Group(GROUP_CONTROL), Group(GROUP_EXPERIMENT)]
        self.events: List[dict] = []

    def to_dict(self) -> dict:
        return {
            "_version": 2,
            "paradigm": self.paradigm,
            "date": self.date,
            "operator": self.operator,
            "remark": self.remark,
            "groups": [g.to_dict() for g in self.groups],
            "events": self.events,
        }

    @classmethod
    def from_dict(cls, d: dict) -> "Session":
        sess = cls.__new__(cls)
        sess.paradigm = d.get("paradigm", PARADIGM_3SIT)
        sess.date = d.get("date", datetime.now().strftime("%Y-%m-%d"))
        sess.operator = d.get("operator", "")
        sess.remark = d.get("remark", "")
        sess.events = d.get("events", [])
        groups_data = d.get("groups", [])
        if groups_data:
            sess.groups = [Group.from_dict(g) for g in groups_data]
            existing_names = {g.name for g in sess.groups}
            for gname in ALL_GROUPS:
                if gname not in existing_names:
                    sess.groups.append(Group(gname))
        else:
            sess.groups = [Group(GROUP_CONTROL), Group(GROUP_EXPERIMENT)]
        return sess

    @classmethod
    def _migrate_from_v1(cls, d: dict) -> "Session":
        """Load from v1 format (single Experiment with items, no groups)."""
        sess = cls()
        paradigm_map = {
            "三箱社交": PARADIGM_3SIT,
            "自由社交": PARADIGM_FREESIT,
            # v1 "都做" combined both paradigms; default to 3-SIT for migration
            "都做": PARADIGM_3SIT,
        }
        sess.paradigm = paradigm_map.get(d.get("paradigm", ""), PARADIGM_3SIT)
        sess.date = d.get("date", datetime.now().strftime("%Y-%m-%d"))
        sess.operator = d.get("operator", "")
        sess.remark = d.get("remark", "")
        sess.events = d.get("events", [])
        old_items = d.get("items", [])
        if old_items:
            name_map = {
                "学习小鼠与真鼠社交的时间": "Mice/s",
                "与玩具社交的时间": "Toy/s",
                "与玩具鼠的社交时间": "Toy/s",
            }
            subj = Subject.__new__(Subject)
            subj.name = d.get("mouse_id", "") or "Subject1"
            subj.paradigm = sess.paradigm
            subj.items = []
            for it_d in old_items:
                it = Item.from_dict(it_d)
                it.name = name_map.get(it.name, it.name)
                subj.items.append(it)
            subj.started = d.get("started", False)
            subj.start_time = d.get("start_time", None)
            sess.groups[0].subjects.append(subj)
        return sess


# ---------------------------------------------------------------------------
# Settings Dialog
# ---------------------------------------------------------------------------

class SettingsDialog(QDialog):
    def __init__(self, settings: dict, parent=None):
        super().__init__(parent)
        self.setWindowTitle("设置")
        self.setMinimumWidth(400)
        self._settings = deepcopy(settings)
        self._build_ui()

    def _build_ui(self):
        layout = QVBoxLayout(self)

        hotkey_group = QGroupBox("快捷键（窗口聚焦时生效）")
        form = QFormLayout()
        hotkey_group.setLayout(form)

        self._hk_start_stop = QKeySequenceEdit(
            QKeySequence(self._settings["hotkeys"].get("start_stop", "F5"))
        )
        self._hk_reset = QKeySequenceEdit(
            QKeySequence(self._settings["hotkeys"].get("reset", "F6"))
        )
        self._hk_export = QKeySequenceEdit(
            QKeySequence(self._settings["hotkeys"].get("export", "F7"))
        )

        form.addRow("开始/暂停:", self._hk_start_stop)
        form.addRow("重置:", self._hk_reset)
        form.addRow("导出:", self._hk_export)
        layout.addWidget(hotkey_group)

        btn_row = QHBoxLayout()
        btn_ok = QPushButton("保存")
        btn_cancel = QPushButton("取消")
        btn_ok.clicked.connect(self._on_ok)
        btn_cancel.clicked.connect(self.reject)
        btn_row.addStretch()
        btn_row.addWidget(btn_ok)
        btn_row.addWidget(btn_cancel)
        layout.addLayout(btn_row)

    def _on_ok(self):
        self._settings["hotkeys"]["start_stop"] = (
            self._hk_start_stop.keySequence().toString()
        )
        self._settings["hotkeys"]["reset"] = (
            self._hk_reset.keySequence().toString()
        )
        self._settings["hotkeys"]["export"] = (
            self._hk_export.keySequence().toString()
        )
        self.accept()

    def get_settings(self) -> dict:
        return self._settings


# ---------------------------------------------------------------------------
# Item row widget (timer or counter)
# ---------------------------------------------------------------------------

class ItemRowWidget(QWidget):
    """A single row representing one timer or counter item."""

    toggled = Signal(int)
    incremented = Signal(int)
    decremented = Signal(int)
    name_changed = Signal(int, str)
    deleted = Signal(int)

    def __init__(self, index: int, item: Item, parent=None):
        super().__init__(parent)
        self._index = index
        self._item = item
        self._build_ui()

    def _build_ui(self):
        row = QHBoxLayout(self)
        row.setContentsMargins(4, 2, 4, 2)

        idx_label = QLabel(f"{self._index + 1}.")
        idx_label.setFixedWidth(24)
        idx_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        font_idx = QFont()
        font_idx.setBold(True)
        idx_label.setFont(font_idx)
        row.addWidget(idx_label)

        self._name_edit = QLineEdit(self._item.name)
        self._name_edit.setMinimumWidth(140)
        self._name_edit.textChanged.connect(self._on_name_changed)
        row.addWidget(self._name_edit)

        self._val_label = QLabel()
        self._val_label.setMinimumWidth(80)
        self._val_label.setAlignment(Qt.AlignCenter)
        font = QFont("Courier New", 12)
        font.setBold(True)
        self._val_label.setFont(font)
        self._update_val_label()
        row.addWidget(self._val_label)

        if self._item.kind == "timer":
            btn = QPushButton("▶/⏸")
            btn.setToolTip("开始/暂停计时")
            btn.setFixedWidth(60)
            btn.clicked.connect(lambda: self.toggled.emit(self._index))
            row.addWidget(btn)
            self._toggle_btn = btn
        else:
            btn_inc = QPushButton("+1")
            btn_inc.setToolTip("计数 +1")
            btn_inc.setFixedWidth(40)
            btn_inc.clicked.connect(lambda: self.incremented.emit(self._index))

            btn_dec = QPushButton("-1")
            btn_dec.setToolTip("计数 -1")
            btn_dec.setFixedWidth(40)
            btn_dec.clicked.connect(lambda: self.decremented.emit(self._index))

            row.addWidget(btn_inc)
            row.addWidget(btn_dec)
            self._toggle_btn = None

        row.addStretch()

        btn_del = QPushButton("✕")
        btn_del.setToolTip("删除此项目")
        btn_del.setFixedWidth(30)
        btn_del.setStyleSheet("color: #cc0000;")
        btn_del.clicked.connect(lambda: self.deleted.emit(self._index))
        row.addWidget(btn_del)

    def _on_name_changed(self, text: str):
        self.name_changed.emit(self._index, text)

    def _update_val_label(self):
        if self._item.kind == "timer":
            self._val_label.setText(fmt_ssxx(self._item.current_elapsed()))
        else:
            self._val_label.setText(str(self._item.count))

    def refresh(self):
        self._update_val_label()
        if self._item.kind == "timer" and self._toggle_btn:
            if self._item.running:
                self._toggle_btn.setStyleSheet("color: red; font-weight: bold;")
            else:
                self._toggle_btn.setStyleSheet("")


# ---------------------------------------------------------------------------
# Subject Detail Dialog
# ---------------------------------------------------------------------------

class SubjectDetailDialog(QDialog):
    """Dialog for timing/counting items for a specific subject."""

    def __init__(
        self,
        subj: Subject,
        group_name: str,
        settings: dict,
        log_callback,
        autosave_callback,
        parent=None,
    ):
        super().__init__(parent)
        self._subj = subj
        self._group_name = group_name
        self._settings = settings
        self._log_callback = log_callback
        self._autosave_callback = autosave_callback
        self._item_row_widgets: List[ItemRowWidget] = []

        self.setWindowTitle(f"{subj.paradigm} — {group_name} — {subj.name}")
        self.setMinimumSize(640, 480)

        self._tick_timer = QTimer(self)
        self._tick_timer.setInterval(100)
        self._tick_timer.timeout.connect(self._on_tick)

        self._build_ui()
        self._tick_timer.start()

    def _build_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(12, 12, 12, 8)
        layout.setSpacing(8)

        title_label = QLabel(
            f"<span style='font-size:14px;font-weight:bold;'>{self._subj.paradigm}</span>"
            f" &nbsp;|&nbsp; {self._group_name} &nbsp;|&nbsp; "
            f"<span style='font-size:14px;font-weight:bold;'>{self._subj.name}</span>"
        )
        title_label.setAlignment(Qt.AlignCenter)
        layout.addWidget(title_label)

        self._status_label = QLabel("就绪")
        self._status_label.setAlignment(Qt.AlignCenter)
        self._status_label.setFixedHeight(28)
        layout.addWidget(self._status_label)

        items_container = QWidget()
        self._items_layout = QVBoxLayout(items_container)
        self._items_layout.setContentsMargins(0, 0, 0, 0)
        self._items_layout.setSpacing(4)

        scroll = QScrollArea()
        scroll.setWidget(items_container)
        scroll.setWidgetResizable(True)
        scroll.setMinimumHeight(180)
        layout.addWidget(scroll)

        item_btn_row = QHBoxLayout()
        btn_add_timer = QPushButton("+ 计时项")
        btn_add_counter = QPushButton("+ 计数项")
        btn_add_timer.clicked.connect(lambda: self._add_item("timer"))
        btn_add_counter.clicked.connect(lambda: self._add_item("counter"))
        item_btn_row.addWidget(btn_add_timer)
        item_btn_row.addWidget(btn_add_counter)
        item_btn_row.addStretch()
        layout.addLayout(item_btn_row)

        ctrl_row = QHBoxLayout()
        btn_back = QPushButton("← 返回汇总表")
        btn_back.setFixedHeight(36)
        btn_back.clicked.connect(self.accept)

        self._btn_start_stop = QPushButton("开始实验 (F5)")
        self._btn_start_stop.setFixedHeight(36)
        self._btn_start_stop.clicked.connect(self._on_start_stop)

        btn_reset = QPushButton("重置 (F6)")
        btn_reset.setFixedHeight(36)
        btn_reset.clicked.connect(self._on_reset)

        ctrl_row.addWidget(btn_back)
        ctrl_row.addStretch()
        ctrl_row.addWidget(self._btn_start_stop)
        ctrl_row.addWidget(btn_reset)
        layout.addLayout(ctrl_row)

        hint = QLabel("快捷键: 数字键 1~9 控制对应项目计时/计数；F5 开始/暂停；F6 重置")
        hint.setStyleSheet("color: gray; font-size: 11px;")
        hint.setAlignment(Qt.AlignCenter)
        layout.addWidget(hint)

        self._rebuild_item_rows()
        self._update_status()

    def _rebuild_item_rows(self):
        while self._items_layout.count():
            child = self._items_layout.takeAt(0)
            if child.widget():
                child.widget().deleteLater()
        self._item_row_widgets.clear()

        for i, item in enumerate(self._subj.items):
            w = ItemRowWidget(i, item)
            w.toggled.connect(self._toggle_timer)
            w.incremented.connect(self._increment_counter)
            w.decremented.connect(self._decrement_counter)
            w.name_changed.connect(self._on_item_name_changed)
            w.deleted.connect(self._delete_item)
            self._items_layout.addWidget(w)
            self._item_row_widgets.append(w)

        self._items_layout.addStretch()

    def _add_item(self, kind: str):
        name = "新计时项" if kind == "timer" else "新计数项"
        self._subj.items.append(Item(kind, name))
        self._rebuild_item_rows()
        self._autosave_callback()

    def _delete_item(self, index: int):
        if 0 <= index < len(self._subj.items):
            self._subj.items.pop(index)
            self._rebuild_item_rows()
            self._autosave_callback()

    def _on_tick(self):
        for w in self._item_row_widgets:
            w.refresh()

    def _update_status(self):
        if self._subj.started:
            self._status_label.setText("⚠ 实验进行中")
            self._status_label.setStyleSheet(
                "background-color: #cc0000; color: white; font-weight: bold;"
            )
            self._btn_start_stop.setText("暂停实验 (F5)")
        else:
            self._status_label.setText("就绪")
            self._status_label.setStyleSheet(
                "background-color: #e8e8e8; color: #555;"
            )
            self._btn_start_stop.setText("开始实验 (F5)")

    def _on_start_stop(self):
        if not self._subj.started:
            self._subj.started = True
            self._subj.start_time = datetime.now().isoformat()
            self._log_callback(
                "experiment_start",
                detail=f"{self._group_name}/{self._subj.name}",
            )
        else:
            any_running = any(
                it.running for it in self._subj.items if it.kind == "timer"
            )
            if any_running:
                for it in self._subj.items:
                    if it.kind == "timer" and it.running:
                        it.elapsed += time.perf_counter() - (
                            it.last_start_ts or time.perf_counter()
                        )
                        it.running = False
                        it.last_start_ts = None
                self._log_callback(
                    "pause_all", detail=f"{self._group_name}/{self._subj.name}"
                )
            else:
                self._log_callback(
                    "resume_all", detail=f"{self._group_name}/{self._subj.name}"
                )
        self._update_status()
        self._autosave_callback()

    def _on_reset(self):
        reply = QMessageBox.question(
            self,
            "确认重置",
            f"确认重置 [{self._subj.name}] 的所有计时/计数数据？",
            QMessageBox.Yes | QMessageBox.No,
        )
        if reply != QMessageBox.Yes:
            return
        self._subj.reset()
        self._rebuild_item_rows()
        self._update_status()
        self._log_callback("reset_all", detail=f"{self._group_name}/{self._subj.name}")
        self._autosave_callback()

    def _toggle_timer(self, index: int):
        if index >= len(self._subj.items):
            return
        item = self._subj.items[index]
        if item.kind != "timer":
            return
        if not self._subj.started:
            QMessageBox.information(self, "提示", '请先点击"开始实验"。')
            return
        now = time.perf_counter()
        if item.running:
            item.elapsed += now - (item.last_start_ts or now)
            item.running = False
            item.last_start_ts = None
            self._log_callback("timer_stop", item_name=item.name)
        else:
            item.running = True
            item.last_start_ts = now
            self._log_callback("timer_start", item_name=item.name)
        self._autosave_callback()

    def _increment_counter(self, index: int):
        if index >= len(self._subj.items):
            return
        item = self._subj.items[index]
        if not self._subj.started:
            QMessageBox.information(self, "提示", '请先点击"开始实验"。')
            return
        item.count += 1
        self._log_callback("counter_inc", item_name=item.name, detail=str(item.count))
        self._autosave_callback()

    def _decrement_counter(self, index: int):
        if index >= len(self._subj.items):
            return
        item = self._subj.items[index]
        if item.count > 0:
            item.count -= 1
        self._log_callback("counter_dec", item_name=item.name, detail=str(item.count))
        self._autosave_callback()

    def _on_item_name_changed(self, index: int, name: str):
        if index < len(self._subj.items):
            self._subj.items[index].name = name

    def keyPressEvent(self, event):
        focused = QApplication.focusWidget()
        if isinstance(focused, (QLineEdit, QTextEdit, QKeySequenceEdit)):
            super().keyPressEvent(event)
            return

        hk = self._settings["hotkeys"]

        def _matches(key_str: str) -> bool:
            if not key_str:
                return False
            qs = QKeySequence(key_str)
            if qs.isEmpty():
                return False
            return event.keyCombination() == qs[0]

        if _matches(hk.get("start_stop", "F5")):
            self._on_start_stop()
            event.accept()
            return
        if _matches(hk.get("reset", "F6")):
            self._on_reset()
            event.accept()
            return

        key = event.key()
        if Qt.Key_1 <= key <= Qt.Key_9:
            idx = key - Qt.Key_1
            if idx < len(self._subj.items):
                item = self._subj.items[idx]
                if item.kind == "timer":
                    self._toggle_timer(idx)
                else:
                    self._increment_counter(idx)
                event.accept()
                return

        super().keyPressEvent(event)

    def closeEvent(self, event):
        self._tick_timer.stop()
        super().closeEvent(event)


# ---------------------------------------------------------------------------
# Summary table widget
# ---------------------------------------------------------------------------

class SummaryTableWidget(QWidget):
    """
    A single QTableWidget showing all groups and their subjects for the
    current paradigm.  Group header rows span all value columns.
    """

    def __init__(
        self,
        session: Session,
        settings_ref: dict,
        log_callback,
        autosave_callback,
        parent=None,
    ):
        super().__init__(parent)
        self._session = session
        self._settings_ref = settings_ref
        self._log_callback = log_callback
        self._autosave_callback = autosave_callback
        self._subject_rows: List[tuple] = []
        self._build_ui()
        self._rebuild()

    def _build_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)

        self._table = QTableWidget()
        self._table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self._table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self._table.verticalHeader().setVisible(False)
        self._table.setAlternatingRowColors(True)
        layout.addWidget(self._table)

    def _col_headers(self) -> List[str]:
        if self._session.paradigm == PARADIGM_3SIT:
            return ["对象名", "Mice/s", "Toy/s", "Mice-Toy/s", "Mice+Toy/s", "DI", "操作"]
        return ["对象名", "适应时间/s", "嗅探时间/s", "躲避次数", "操作"]

    def _rebuild(self):
        self._subject_rows = []
        headers = self._col_headers()
        ncols = len(headers)

        self._table.clearContents()
        self._table.setColumnCount(ncols)
        self._table.setHorizontalHeaderLabels(headers)

        hh = self._table.horizontalHeader()
        for i in range(ncols - 1):
            hh.setSectionResizeMode(i, QHeaderView.Stretch)
        hh.setSectionResizeMode(ncols - 1, QHeaderView.Fixed)
        self._table.setColumnWidth(ncols - 1, 130)

        total_rows = sum(
            1 + len(g.subjects)
            for g in self._session.groups
        )
        self._table.setRowCount(total_rows)

        current_row = 0
        for g_idx, group in enumerate(self._session.groups):
            # Group header row
            self._table.setSpan(current_row, 0, 1, ncols - 1)
            header_item = QTableWidgetItem(f"  {group.name}")
            header_item.setTextAlignment(Qt.AlignVCenter | Qt.AlignLeft)
            bg_color = QColor("#4a6741") if g_idx == 0 else QColor("#3a5978")
            header_item.setBackground(bg_color)
            header_item.setForeground(QColor("white"))
            hdr_font = QFont()
            hdr_font.setBold(True)
            hdr_font.setPointSize(10)
            header_item.setFont(hdr_font)
            self._table.setItem(current_row, 0, header_item)
            self._table.setRowHeight(current_row, 34)

            btn_add = QPushButton("+ 添加对象")
            btn_add.setStyleSheet(
                "color: white; background-color: #5a8050; "
                "border-radius: 3px; padding: 2px 6px;"
            )
            btn_add.clicked.connect(
                lambda _, gi=g_idx: self._add_subject(gi)
            )
            self._table.setCellWidget(current_row, ncols - 1, btn_add)
            current_row += 1

            for s_idx, subj in enumerate(group.subjects):
                self._subject_rows.append((current_row, g_idx, s_idx))
                self._fill_subject_row(current_row, g_idx, s_idx, ncols)
                self._table.setRowHeight(current_row, 30)
                current_row += 1

    def _fill_subject_row(self, row: int, g_idx: int, s_idx: int, ncols: int):
        subj = self._session.groups[g_idx].subjects[s_idx]

        self._table.setItem(row, 0, QTableWidgetItem(subj.name))

        if self._session.paradigm == PARADIGM_3SIT:
            mice = subj.get_value("Mice/s")
            toy = subj.get_value("Toy/s")
            mice_toy = mice - toy

            self._table.setItem(row, 1, QTableWidgetItem(fmt_ssxx(mice)))
            self._table.setItem(row, 2, QTableWidgetItem(fmt_ssxx(toy)))
            self._table.setItem(row, 3, QTableWidgetItem(fmt_ssxx_signed(mice_toy)))
            self._table.setItem(row, 4, QTableWidgetItem(fmt_ssxx(mice + toy)))
            self._table.setItem(row, 5, QTableWidgetItem(calc_di(mice, toy)))
        else:
            adapt = subj.get_value("适应时间/s")
            sniff = subj.get_value("嗅探时间/s")
            avoid = subj.get_value("躲避次数")

            self._table.setItem(row, 1, QTableWidgetItem(fmt_ssxx(adapt)))
            self._table.setItem(row, 2, QTableWidgetItem(fmt_ssxx(sniff)))
            self._table.setItem(row, 3, QTableWidgetItem(str(int(avoid))))

        op_w = self._make_op_widget(g_idx, s_idx)
        self._table.setCellWidget(row, ncols - 1, op_w)

    def _make_op_widget(self, g_idx: int, s_idx: int) -> QWidget:
        w = QWidget()
        lay = QHBoxLayout(w)
        lay.setContentsMargins(3, 2, 3, 2)
        lay.setSpacing(4)

        btn_detail = QPushButton("详情")
        btn_detail.setFixedHeight(24)
        btn_del = QPushButton("删除")
        btn_del.setFixedHeight(24)
        btn_del.setStyleSheet("color: #cc0000;")

        btn_detail.clicked.connect(
            lambda _, gi=g_idx, si=s_idx: self._open_detail(gi, si)
        )
        btn_del.clicked.connect(
            lambda _, gi=g_idx, si=s_idx: self._delete_subject(gi, si)
        )

        lay.addWidget(btn_detail)
        lay.addWidget(btn_del)
        return w

    def _add_subject(self, g_idx: int):
        group = self._session.groups[g_idx]
        name, ok = QInputDialog.getText(
            self, f"添加实验对象 — {group.name}", "对象名称（例如 B21）:"
        )
        if not ok or not name.strip():
            return
        name = name.strip()
        if any(s.name == name for s in group.subjects):
            QMessageBox.warning(
                self, "重名", f"实验对象 [{name}] 已存在于 {group.name} 中。"
            )
            return
        group.subjects.append(Subject(name, self._session.paradigm))
        self._rebuild()
        self._autosave_callback()

    def _delete_subject(self, g_idx: int, s_idx: int):
        group = self._session.groups[g_idx]
        if s_idx >= len(group.subjects):
            return
        subj = group.subjects[s_idx]
        reply = QMessageBox.question(
            self,
            "确认删除",
            f"确认删除实验对象 [{subj.name}]？\n该对象的所有计时/计数数据将被清除，且无法恢复。",
            QMessageBox.Yes | QMessageBox.No,
        )
        if reply != QMessageBox.Yes:
            return
        group.subjects.pop(s_idx)
        self._rebuild()
        self._autosave_callback()

    def _open_detail(self, g_idx: int, s_idx: int):
        group = self._session.groups[g_idx]
        if s_idx >= len(group.subjects):
            return
        subj = group.subjects[s_idx]
        dlg = SubjectDetailDialog(
            subj=subj,
            group_name=group.name,
            settings=self._settings_ref,
            log_callback=self._log_callback,
            autosave_callback=self._autosave_callback,
            parent=self,
        )
        dlg.exec()
        self.refresh()

    def refresh(self):
        """Update only value cells without rebuilding the table structure."""
        for table_row, g_idx, s_idx in self._subject_rows:
            subj = self._session.groups[g_idx].subjects[s_idx]
            if self._session.paradigm == PARADIGM_3SIT:
                mice = subj.get_value("Mice/s")
                toy = subj.get_value("Toy/s")
                mice_toy = mice - toy

                for col, txt in enumerate(
                    [
                        fmt_ssxx(mice),
                        fmt_ssxx(toy),
                        fmt_ssxx_signed(mice_toy),
                        fmt_ssxx(mice + toy),
                        calc_di(mice, toy),
                    ],
                    start=1,
                ):
                    item = self._table.item(table_row, col)
                    if item:
                        item.setText(txt)
            else:
                adapt = subj.get_value("适应时间/s")
                sniff = subj.get_value("嗅探时间/s")
                avoid = subj.get_value("躲避次数")
                for col, txt in enumerate(
                    [fmt_ssxx(adapt), fmt_ssxx(sniff), str(int(avoid))],
                    start=1,
                ):
                    item = self._table.item(table_row, col)
                    if item:
                        item.setText(txt)


# ---------------------------------------------------------------------------
# Main Window
# ---------------------------------------------------------------------------

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("MiceTimer v2.0 - 实验计时计数软件")
        self.setMinimumSize(QSize(900, 620))

        self._settings = load_settings()
        self._session = Session()
        self._summary_widget: Optional[SummaryTableWidget] = None

        self._tick_timer = QTimer(self)
        self._tick_timer.setInterval(100)
        self._tick_timer.timeout.connect(self._on_tick)

        self._build_ui()
        self._try_recover()
        self._tick_timer.start()

    # ------------------------------------------------------------------
    # UI construction
    # ------------------------------------------------------------------

    def _build_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        main_layout = QVBoxLayout(central)
        main_layout.setContentsMargins(8, 8, 8, 8)

        self._status_bar = QStatusBar()
        self.setStatusBar(self._status_bar)
        self._status_label = QLabel("就绪")
        self._status_bar.addWidget(self._status_label)

        tabs = QTabWidget()
        self._tabs = tabs
        main_layout.addWidget(tabs)

        # --- Tab: 汇总 ---
        summary_tab = QWidget()
        summary_layout = QVBoxLayout(summary_tab)
        summary_layout.setContentsMargins(8, 8, 8, 8)
        tabs.addTab(summary_tab, "汇总")

        info_row = QHBoxLayout()
        info_row.addWidget(QLabel("范式:"))
        self._paradigm_combo = QComboBox()
        self._paradigm_combo.addItems(ALL_PARADIGMS)
        self._paradigm_combo.setFixedWidth(110)
        self._paradigm_combo.currentTextChanged.connect(self._on_paradigm_changed)
        info_row.addWidget(self._paradigm_combo)

        info_row.addSpacing(16)
        info_row.addWidget(QLabel("日期:"))
        self._field_date = QLineEdit()
        self._field_date.setFixedWidth(100)
        self._field_date.textChanged.connect(lambda v: setattr(self._session, "date", v))
        info_row.addWidget(self._field_date)

        info_row.addSpacing(8)
        info_row.addWidget(QLabel("实验员:"))
        self._field_operator = QLineEdit()
        self._field_operator.setFixedWidth(100)
        self._field_operator.textChanged.connect(
            lambda v: setattr(self._session, "operator", v)
        )
        info_row.addWidget(self._field_operator)

        info_row.addSpacing(8)
        info_row.addWidget(QLabel("备注:"))
        self._field_remark = QLineEdit()
        self._field_remark.setMinimumWidth(120)
        self._field_remark.textChanged.connect(
            lambda v: setattr(self._session, "remark", v)
        )
        info_row.addWidget(self._field_remark)
        info_row.addStretch()
        summary_layout.addLayout(info_row)

        self._summary_scroll = QScrollArea()
        self._summary_scroll.setWidgetResizable(True)
        summary_layout.addWidget(self._summary_scroll, stretch=1)

        export_row = QHBoxLayout()
        hk = self._settings["hotkeys"]
        self._btn_export_default = QPushButton(
            f"导出到默认目录 ({hk.get('export', 'F7')})"
        )
        self._btn_export_default.clicked.connect(self.export_excel_default)
        self._btn_export_as = QPushButton("另存为...")
        self._btn_export_as.clicked.connect(self.export_excel_as)
        export_row.addWidget(self._btn_export_default)
        export_row.addWidget(self._btn_export_as)
        export_row.addStretch()
        summary_layout.addLayout(export_row)

        # --- Tab: 模板 ---
        tpl_tab = QWidget()
        tpl_layout = QVBoxLayout(tpl_tab)
        tabs.addTab(tpl_tab, "模板")

        tpl_btn_row = QHBoxLayout()
        btn_save_tpl = QPushButton("保存当前会话为模板...")
        btn_load_tpl = QPushButton("加载模板文件...")
        btn_save_tpl.clicked.connect(self._save_template)
        btn_load_tpl.clicked.connect(self._load_template_dialog)
        tpl_btn_row.addWidget(btn_save_tpl)
        tpl_btn_row.addWidget(btn_load_tpl)
        tpl_btn_row.addStretch()
        tpl_layout.addLayout(tpl_btn_row)

        recent_group = QGroupBox("最近模板（快速加载）")
        recent_layout = QVBoxLayout()
        recent_group.setLayout(recent_layout)
        self._recent_list_widget = QTableWidget()
        self._recent_list_widget.setColumnCount(3)
        self._recent_list_widget.setHorizontalHeaderLabels(["模板名", "范式", "对象数"])
        self._recent_list_widget.horizontalHeader().setSectionResizeMode(
            0, QHeaderView.Stretch
        )
        self._recent_list_widget.setSelectionBehavior(QAbstractItemView.SelectRows)
        self._recent_list_widget.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self._recent_list_widget.setMinimumHeight(140)
        recent_layout.addWidget(self._recent_list_widget)

        btn_load_recent = QPushButton("加载选中模板")
        btn_load_recent.clicked.connect(self._load_recent_selected)
        recent_layout.addWidget(btn_load_recent)
        tpl_layout.addWidget(recent_group)
        tpl_layout.addStretch()
        self._refresh_recent_list()

        # --- Tab: 事件日志 ---
        log_tab = QWidget()
        log_layout = QVBoxLayout(log_tab)
        tabs.addTab(log_tab, "事件日志")

        self._log_text = QTextEdit()
        self._log_text.setReadOnly(True)
        log_layout.addWidget(self._log_text)

        btn_clear_log = QPushButton("清空日志")
        btn_clear_log.clicked.connect(self._clear_log)
        log_layout.addWidget(btn_clear_log)

        # --- Tab: 设置 ---
        settings_tab = QWidget()
        settings_layout = QVBoxLayout(settings_tab)
        tabs.addTab(settings_tab, "设置")

        btn_open_settings = QPushButton("打开快捷键设置...")
        btn_open_settings.clicked.connect(self._open_settings_dialog)
        settings_layout.addWidget(btn_open_settings)

        self._hk_info_label = QLabel("当前快捷键：\n" + self._hotkey_summary())
        settings_layout.addWidget(self._hk_info_label)
        settings_layout.addStretch()

    def _build_summary_widget(self):
        w = SummaryTableWidget(
            session=self._session,
            settings_ref=self._settings,
            log_callback=self._log_event,
            autosave_callback=self.save_autosave,
        )
        self._summary_scroll.setWidget(w)
        self._summary_widget = w

    # ------------------------------------------------------------------
    # Paradigm selector
    # ------------------------------------------------------------------

    def _on_paradigm_changed(self, new_paradigm: str):
        if new_paradigm == self._session.paradigm:
            return
        has_subjects = any(len(g.subjects) > 0 for g in self._session.groups)
        if has_subjects:
            reply = QMessageBox.question(
                self,
                "切换范式",
                f"切换到 [{new_paradigm}] 将清除所有实验对象及其数据。\n确认继续？",
                QMessageBox.Yes | QMessageBox.No,
            )
            if reply != QMessageBox.Yes:
                self._paradigm_combo.blockSignals(True)
                self._paradigm_combo.setCurrentText(self._session.paradigm)
                self._paradigm_combo.blockSignals(False)
                return

        self._session.paradigm = new_paradigm
        for group in self._session.groups:
            group.subjects.clear()
        self._build_summary_widget()
        self.save_autosave()

    # ------------------------------------------------------------------
    # Hotkeys
    # ------------------------------------------------------------------

    def _hotkey_summary(self) -> str:
        hk = self._settings["hotkeys"]
        return (
            f"  开始/暂停: {hk.get('start_stop', 'F5')}\n"
            f"  重置: {hk.get('reset', 'F6')}\n"
            f"  导出: {hk.get('export', 'F7')}"
        )

    def keyPressEvent(self, event):
        focused = QApplication.focusWidget()
        if isinstance(focused, (QLineEdit, QTextEdit, QKeySequenceEdit)):
            super().keyPressEvent(event)
            return

        hk = self._settings["hotkeys"]

        def _matches(key_str: str) -> bool:
            if not key_str:
                return False
            qs = QKeySequence(key_str)
            if qs.isEmpty():
                return False
            return event.keyCombination() == qs[0]

        if _matches(hk.get("export", "F7")):
            self.export_excel_default()
            event.accept()
            return

        super().keyPressEvent(event)

    # ------------------------------------------------------------------
    # Tick (UI refresh)
    # ------------------------------------------------------------------

    def _on_tick(self):
        if self._summary_widget:
            self._summary_widget.refresh()

    # ------------------------------------------------------------------
    # Autosave / Recovery
    # ------------------------------------------------------------------

    def save_autosave(self):
        try:
            data = self._session.to_dict()
            with open(RECOVERY_FILE, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
        except Exception:
            pass

    def _try_recover(self):
        if not os.path.exists(RECOVERY_FILE):
            self._apply_session(Session())
            return
        try:
            with open(RECOVERY_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
            if data.get("_version", 1) >= 2:
                sess = Session.from_dict(data)
            else:
                sess = Session._migrate_from_v1(data)

            reply = QMessageBox.question(
                self,
                "恢复上次会话",
                "检测到上次未保存的会话数据，是否恢复？",
                QMessageBox.Yes | QMessageBox.No,
            )
            if reply == QMessageBox.Yes:
                self._apply_session(sess)
            else:
                self._apply_session(Session())
        except Exception:
            self._apply_session(Session())

    def _apply_session(self, sess: Session):
        self._session = sess
        self._sync_fields()
        self._build_summary_widget()

    def _sync_fields(self):
        for field, attr in [
            (self._field_date, "date"),
            (self._field_operator, "operator"),
            (self._field_remark, "remark"),
        ]:
            field.blockSignals(True)
            field.setText(getattr(self._session, attr))
            field.blockSignals(False)

        self._paradigm_combo.blockSignals(True)
        self._paradigm_combo.setCurrentText(self._session.paradigm)
        self._paradigm_combo.blockSignals(False)

    # ------------------------------------------------------------------
    # Event log
    # ------------------------------------------------------------------

    def _log_event(self, action: str, item_name: str = "", detail: str = ""):
        ts = datetime.now().strftime("%H:%M:%S.%f")[:-3]
        entry = {"ts": ts, "action": action, "item": item_name, "detail": detail}
        self._session.events.append(entry)
        line = f"[{ts}] {action}"
        if item_name:
            line += f" | {item_name}"
        if detail:
            line += f" | {detail}"
        self._log_text.append(line)

    def _clear_log(self):
        self._session.events.clear()
        self._log_text.clear()

    # ------------------------------------------------------------------
    # Templates
    # ------------------------------------------------------------------

    def _save_template(self):
        name, ok = QInputDialog.getText(
            self, "保存模板", "模板名称:", text=self._session.paradigm or "模板"
        )
        if not ok or not name.strip():
            return
        name = name.strip()
        template = {
            "name": name,
            "_version": 2,
            "session": self._session.to_dict(),
        }
        tpl_path = os.path.join(TEMPLATES_DIR, name + ".json")
        with open(tpl_path, "w", encoding="utf-8") as f:
            json.dump(template, f, ensure_ascii=False, indent=2)
        recent_entry = {
            "name": name,
            "paradigm": self._session.paradigm,
            "subject_count": sum(len(g.subjects) for g in self._session.groups),
            "_version": 2,
            "session": self._session.to_dict(),
        }
        save_recent_template(recent_entry)
        self._refresh_recent_list()
        QMessageBox.information(self, "保存成功", f"模板 [{name}] 已保存。")

    def _load_template_dialog(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "加载模板", TEMPLATES_DIR, "JSON Files (*.json)"
        )
        if not path:
            return
        try:
            with open(path, "r", encoding="utf-8") as f:
                tpl = json.load(f)
            self._apply_template(tpl)
        except Exception as e:
            QMessageBox.warning(self, "加载失败", str(e))

    def _apply_template(self, tpl: dict):
        try:
            if tpl.get("_version", 1) >= 2 and "session" in tpl:
                sess = Session.from_dict(tpl["session"])
            else:
                sess = Session._migrate_from_v1(tpl)
            self._apply_session(sess)
            self.save_autosave()
            QMessageBox.information(
                self, "已加载", f"模板 [{tpl.get('name', '')}] 已加载。"
            )
        except Exception as e:
            QMessageBox.warning(self, "加载失败", str(e))

    def _refresh_recent_list(self):
        recent = load_recent_templates()
        self._recent_list_widget.setRowCount(len(recent))
        for i, tpl in enumerate(recent):
            self._recent_list_widget.setItem(i, 0, QTableWidgetItem(tpl.get("name", "")))
            self._recent_list_widget.setItem(
                i, 1, QTableWidgetItem(tpl.get("paradigm", ""))
            )
            if "session" in tpl:
                n = sum(
                    len(g.get("subjects", []))
                    for g in tpl["session"].get("groups", [])
                )
            else:
                n = tpl.get("subject_count", 0)
            self._recent_list_widget.setItem(i, 2, QTableWidgetItem(str(n)))

    def _load_recent_selected(self):
        row = self._recent_list_widget.currentRow()
        if row < 0:
            QMessageBox.information(self, "提示", "请先选择一个模板。")
            return
        recent = load_recent_templates()
        if row >= len(recent):
            return
        self._apply_template(recent[row])
        self._refresh_recent_list()

    # ------------------------------------------------------------------
    # Settings
    # ------------------------------------------------------------------

    def _open_settings_dialog(self):
        dlg = SettingsDialog(self._settings, self)
        if dlg.exec() == QDialog.Accepted:
            self._settings = dlg.get_settings()
            save_settings(self._settings)
            self._hk_info_label.setText("当前快捷键：\n" + self._hotkey_summary())
            hk = self._settings["hotkeys"]
            self._btn_export_default.setText(
                f"导出到默认目录 ({hk.get('export', 'F7')})"
            )

    # ------------------------------------------------------------------
    # Excel export
    # ------------------------------------------------------------------

    def _build_export_filename(self) -> str:
        date = safe_name_part(self._session.date) or datetime.now().strftime("%Y-%m-%d")
        paradigm = safe_name_part(self._session.paradigm)
        operator = safe_name_part(self._session.operator)
        parts = [date, paradigm]
        if operator:
            parts.append(operator)
        return "_".join(parts) + ".xlsx"

    def _freeze_all_timers(self):
        for group in self._session.groups:
            for subj in group.subjects:
                for it in subj.items:
                    if it.kind == "timer" and it.running and it.last_start_ts is not None:
                        it.elapsed += time.perf_counter() - it.last_start_ts
                        it.running = False
                        it.last_start_ts = None

    def _write_excel_to_path(self, path: str):
        wb = Workbook()
        ws = wb.active
        ws.title = self._session.paradigm

        paradigm = self._session.paradigm

        # Derive batch prefix from remark (if set) or date
        batch_prefix = self._session.remark.strip() if self._session.remark.strip() else self._session.date
        batch_title = f"{batch_prefix}-{paradigm}" if batch_prefix else paradigm

        # Column definitions per paradigm
        if paradigm == PARADIGM_3SIT:
            col_headers = ["", "Mice/s", "Toy/s", "Mice-Toy/s", "Mice+Toy/s", "DI"]
            col_widths = [14, 12, 12, 14, 14, 16]
            di_col = col_headers.index("DI") + 1  # 1-indexed column number for DI
        else:
            col_headers = ["", "适应时间/s", "嗅探时间/s", "躲避次数"]
            col_widths = [14, 14, 14, 12]
            di_col = None

        ncols = len(col_headers)

        # Set column widths
        for c_idx, w in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(c_idx)].width = w

        # Style helpers
        font_title = Font(bold=True, size=12)
        font_header = Font(bold=True, size=10)
        font_group = Font(bold=True, color="FFFFFF", size=10)
        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center", indent=1)
        fill_control = PatternFill("solid", fgColor="4A6741")
        fill_experiment = PatternFill("solid", fgColor="3A5978")

        row = 1

        # --- Title row (merged, centered) ---
        title_cell = ws.cell(row=row, column=1, value=batch_title)
        ws.merge_cells(
            start_row=row, start_column=1, end_row=row, end_column=ncols
        )
        title_cell.font = font_title
        title_cell.alignment = align_center
        ws.row_dimensions[row].height = 22
        row += 1

        # --- Header row ---
        for c_idx, hdr in enumerate(col_headers, start=1):
            cell = ws.cell(row=row, column=c_idx, value=hdr)
            cell.font = font_header
            cell.alignment = align_center
        ws.row_dimensions[row].height = 18
        row += 1

        # --- Group blocks ---
        for group in self._session.groups:
            # Group title row (merged)
            group_cell = ws.cell(row=row, column=1, value=group.name)
            ws.merge_cells(
                start_row=row, start_column=1, end_row=row, end_column=ncols
            )
            group_fill = fill_control if group.name == GROUP_CONTROL else fill_experiment
            group_cell.fill = group_fill
            group_cell.font = font_group
            group_cell.alignment = align_left
            ws.row_dimensions[row].height = 18
            row += 1

            # Subject rows
            for subj in group.subjects:
                name_cell = ws.cell(row=row, column=1, value=subj.name)
                name_cell.alignment = align_center

                if paradigm == PARADIGM_3SIT:
                    mice = subj.get_value("Mice/s")
                    toy = subj.get_value("Toy/s")
                    mice_toy = mice - toy
                    mice_plus_toy = mice + toy

                    for c_idx, (val, fmt) in enumerate(
                        [
                            (round(mice, 2), "0.00"),
                            (round(toy, 2), "0.00"),
                            (round(mice_toy, 2), "0.00"),
                            (round(mice_plus_toy, 2), "0.00"),
                        ],
                        start=2,
                    ):
                        cell = ws.cell(row=row, column=c_idx, value=val)
                        cell.number_format = fmt
                        cell.alignment = align_center

                    di_val = calc_di_numeric(mice, toy)
                    di_cell = ws.cell(
                        row=row, column=di_col, value=di_val if di_val is not None else ""
                    )
                    if di_val is not None:
                        di_cell.number_format = "0.000000"
                    di_cell.alignment = align_center

                else:  # Free-SIT
                    adapt = subj.get_value("适应时间/s")
                    sniff = subj.get_value("嗅探时间/s")
                    avoid = int(subj.get_value("躲避次数"))

                    for c_idx, (val, fmt) in enumerate(
                        [
                            (round(adapt, 2), "0.00"),
                            (round(sniff, 2), "0.00"),
                        ],
                        start=2,
                    ):
                        cell = ws.cell(row=row, column=c_idx, value=val)
                        cell.number_format = fmt
                        cell.alignment = align_center

                    avoid_cell = ws.cell(row=row, column=4, value=avoid)
                    avoid_cell.alignment = align_center

                ws.row_dimensions[row].height = 16
                row += 1

        # Blank separator row after block
        row += 1

        # --- Events sheet ---
        ws2 = wb.create_sheet("事件日志")
        ws2.append(["时间戳", "动作", "项目", "详情"])
        for e in self._session.events:
            ws2.append([e["ts"], e["action"], e.get("item", ""), e.get("detail", "")])

        wb.save(path)
        QApplication.beep()
        self._log_event("export_excel", detail=path)
        self.save_autosave()

    def export_excel_default(self):
        self._freeze_all_timers()
        fname = self._build_export_filename()
        path = os.path.join(EXPORT_DIR, fname)
        base, ext = os.path.splitext(path)
        idx = 1
        while os.path.exists(path):
            path = f"{base}_{idx}{ext}"
            idx += 1
        try:
            self._write_excel_to_path(path)
            QMessageBox.information(self, "导出成功", f"已导出到默认目录：\n{path}")
        except Exception as e:
            QMessageBox.warning(self, "导出失败", str(e))

    def export_excel_as(self):
        self._freeze_all_timers()
        fname = self._build_export_filename()
        default_path = os.path.join(EXPORT_DIR, fname)
        path, _ = QFileDialog.getSaveFileName(
            self, "另存为", default_path, "Excel Files (*.xlsx)"
        )
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"
        try:
            self._write_excel_to_path(path)
            QMessageBox.information(self, "导出成功", f"已导出：\n{path}")
        except Exception as e:
            QMessageBox.warning(self, "导出失败", str(e))

    # ------------------------------------------------------------------
    # Close event
    # ------------------------------------------------------------------

    def closeEvent(self, event):
        self.save_autosave()
        self._tick_timer.stop()
        event.accept()


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def main():
    app = QApplication(sys.argv)
    app.setApplicationName("MiceTimer")
    app.setApplicationVersion("2.0.0")

    window = MainWindow()
    window.show()

    sys.exit(app.exec())


if __name__ == "__main__":
    main()
